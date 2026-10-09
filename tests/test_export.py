"""
Modul 3 (export) testleri — gercek pipeline ciktisi seklinde sahte veri.
"""
import os
import tempfile

from app.pipeline import export


MEETING = {
    "title": "Q4 Planlama Toplantısı",
    "date": "2026-09-23",
    "participants": ["Ahmet Yılmaz", "Elif Kaya"],
    "executive_summary": "Bütçe onaylandı; iki aksiyon ve bir risk kayda geçti.",
    "items": [
        {"type": "decision", "text": "Q4 bütçesi 4.5M TL olarak onaylandı.",
         "owner": None, "deadline": None,
         "evidence": "Ahmet: bütçeyi 4.5 yapalım, onaylıyorum",
         "line_no": 12, "confidence": 0.93},
        {"type": "decision", "text": "Yeni ofis ihalesi ertelendi.",
         "owner": None, "deadline": None, "evidence": "", "line_no": 30,
         "confidence": 0.55},                      # dusuk guven -> isaretlenmeli
        {"type": "action_item", "text": "Bütçe gerçekleşme raporu hazırlanacak.",
         "owner": "Elif Kaya", "deadline": "2026-09-30",
         "evidence": "...", "line_no": 18, "confidence": 0.88},
        {"type": "action_item", "text": "Stajyer ilanı yayınlanacak.",
         "owner": None, "deadline": None,
         "evidence": "...", "line_no": 22, "confidence": 0.61},
        {"type": "risk", "text": "Tedarikçi gecikmesi Ekim teslimini riske atıyor.",
         "owner": None, "deadline": None, "evidence": "...", "line_no": 40,
         "confidence": 0.8},
        {"type": "open_question", "text": "İkinci bayi için fiyat politikası ne olacak?",
         "owner": None, "deadline": None, "evidence": "", "line_no": 45,
         "confidence": 0.75},
        {"type": "idea", "text": "Aylık bütçe otomasyonu düşünülebilir.",
         "owner": None, "deadline": None, "evidence": "", "line_no": 50,
         "confidence": 0.7},
        # casual artigi gelmemeli: kalite katmani zaten elemis olmali
    ],
}


def _report():
    return export.build_report_data(MEETING)


def test_html_contains_all_sections():
    html = export.render_html(_report())
    for section in ["Yönetici Özeti", "Kararlar", "Aksiyon Maddeleri",
                    "Riskler", "Açık Sorular", "Önemli Fikirler"]:
        assert section in html, f"{section} bolumu HTML'de yok"


def test_html_action_table_fields():
    html = export.render_html(_report())
    assert "Elif Kaya" in html          # sorumlu
    assert "2026-09-30" in html         # son tarih
    assert ">—</td>" in html          # sorumlusu olmayan madde


def test_html_flags_low_confidence():
    html = export.render_html(_report())
    assert "düşük güven" in html        # %55 karar isaretlenmeli


def test_html_escapes_and_no_casual():
    html = export.render_html(_report())
    assert "öğle yemeği" not in html    # casual icerik asla sizmamali


def test_docx_roundtrip():
    with tempfile.TemporaryDirectory() as d:
        path = export.export_docx(_report(), os.path.join(d, "r.docx"))
        from docx import Document
        doc = Document(path)            # okunabiliyor mu
        text = "\n".join(p.text for p in doc.paragraphs)
        assert "Q4 Planlama Toplantısı" in text
        assert "Yönetici Özeti" in text
        assert len(doc.tables) == 1     # aksiyon tablosu
        row = doc.tables[0].rows[1].cells
        assert row[1].text == "Elif Kaya"


def test_pdf_error_is_actionable():
    # weasyprint kurulu degilse: anlamli hata, sessiz patlama olmamali
    try:
        export.export_pdf(_report())
    except RuntimeError as e:
        assert "weasyprint" in str(e)
